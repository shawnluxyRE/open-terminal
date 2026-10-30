"""REQ-002 auxiliary inputs must keep the caller's file-access boundary."""

import io
import json
import os
from pathlib import Path
import pwd
import shutil
import subprocess
import sys
import tempfile
import uuid
import zipfile

from docx import Document
from docx.shared import Inches
from lxml import etree
from PIL import Image
from pypdf import PdfReader

from test_file_http import denied


PRIVATE_COLOR = (243, 17, 99)


def _normal_document():
    placeholder = io.BytesIO()
    Image.new("RGB", (100, 100), (255, 255, 255)).save(placeholder, format="PNG")
    placeholder.seek(0)
    document = Document()
    document.add_paragraph("Alice's document with a linked picture")
    document.add_picture(placeholder, width=Inches(1))
    original = io.BytesIO()
    document.save(original)
    return original.getvalue()


def _linked_document(original, secret):
    # Word's ordinary external-picture representation: r:link instead of
    # r:embed, with an External relationship pointing to a file URI.
    linked = io.BytesIO()
    relationships = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(
        linked, "w", zipfile.ZIP_DEFLATED
    ) as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "word/_rels/document.xml.rels":
                xml = etree.fromstring(content)
                for relationship in xml:
                    if relationship.get("Type", "").endswith("/image"):
                        relationship.set("Target", secret.as_uri())
                        relationship.set("TargetMode", "External")
                content = etree.tostring(xml)
            elif info.filename == "word/document.xml":
                xml = etree.fromstring(content)
                for picture in xml.iter(
                    "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
                ):
                    reference = picture.attrib.pop("{" + relationships + "}embed")
                    picture.set("{" + relationships + "}link", reference)
                content = etree.tostring(xml)
            target.writestr(info, content)
    return linked.getvalue()


def _preview_colors(response):
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/pdf")
    return [
        picture.image.convert("RGB").getpixel((0, 0))
        for page in PdfReader(io.BytesIO(response.content)).pages
        for picture in page.images
    ]


def _assert_no_private_picture(response, linked):
    if response.status_code == 403:
        denied(response)
        return
    assert response.status_code == 200, response.text
    if not response.headers["content-type"].startswith("application/pdf"):
        # Returning the authorized source without resolving external content
        # does not expose the other account's image bytes.
        assert response.content == linked
        return
    assert PRIVATE_COLOR not in _preview_colors(response), (
        "Office preview embedded Bob's private linked image although the "
        "same source is denied by the direct file API"
    )


def test_office_preview_does_not_embed_another_users_linked_image(api):
    """Static external DOCX pictures must keep the caller's file boundary.

    The real converter must reject an unauthorized dependency or render without
    its private pixels. No script or concurrent path change is involved.
    """
    assert shutil.which("soffice") or shutil.which("libreoffice")
    secret = api.homes["bob"] / "private-photo.png"
    Image.new("RGB", (100, 100), PRIVATE_COLOR).save(secret)
    denied(api.request("GET", "/files/view", params={"path": str(secret)}))
    original = _normal_document()
    normal = api.homes["alice"] / "embedded-picture.docx"
    normal.write_bytes(original)
    positive = api.request(
        "GET", "/files/view", params={"path": str(normal), "preview": True}
    )
    assert (255, 255, 255) in _preview_colors(positive)

    linked = _linked_document(original, secret)
    owned = api.homes["alice"] / "linked-picture.docx"
    owned.write_bytes(linked)

    response = api.request(
        "GET", "/files/view", params={"path": str(owned), "preview": True}
    )
    _assert_no_private_picture(response, linked)
    assert secret.exists()
    assert Image.open(secret).getpixel((0, 0)) == PRIVATE_COLOR


def _service_worker():
    """Repeat through real identity provisioning under the actual service UID."""
    from fastapi.testclient import TestClient
    from open_terminal import main

    assert os.geteuid() == 1000
    marker = uuid.uuid4().hex
    users = {name: "req002-office-reference-" + name + "-" + marker
             for name in ("alice", "bob")}
    with TestClient(main.app) as client:
        def request(user, method, endpoint, **kwargs):
            return client.request(method, endpoint, headers={
                "Authorization": "Bearer req002-auxiliary-test-not-a-secret",
                "X-User-Id": users[user], "X-Session-Id": "office-reference",
            }, **kwargs)

        homes = {}
        for user in users:
            result = request(user, "GET", "/files/cwd")
            assert result.status_code == 200, result.text
            homes[user] = Path(result.json()["home"])

        def upload(user, name, data):
            result = request(user, "POST", "/files/upload",
                             params={"directory": str(homes[user])},
                             files={"file": (name, data)})
            assert result.status_code == 200, result.text

        png = io.BytesIO()
        Image.new("RGB", (100, 100), PRIVATE_COLOR).save(png, format="PNG")
        upload("bob", "private-photo.png", png.getvalue())
        secret = homes["bob"] / "private-photo.png"
        owner = secret.stat()
        bob = pwd.getpwnam(homes["bob"].name)
        assert owner.st_uid == bob.pw_uid and owner.st_uid != 1000
        assert owner.st_gid == bob.pw_gid and owner.st_gid in os.getgroups()
        direct = request("alice", "GET", "/files/view", params={"path": str(secret)})
        denied(direct)

        original = _normal_document()
        upload("alice", "embedded-picture.docx", original)
        positive = request("alice", "GET", "/files/view", params={
            "path": str(homes["alice"] / "embedded-picture.docx"), "preview": True,
        })
        assert (255, 255, 255) in _preview_colors(positive)
        linked = _linked_document(original, secret)
        upload("alice", "linked-picture.docx", linked)
        response = request("alice", "GET", "/files/view", params={
            "path": str(homes["alice"] / "linked-picture.docx"), "preview": True,
        })
        private_pixels = (response.status_code == 200
                          and response.headers["content-type"].startswith("application/pdf")
                          and PRIVATE_COLOR in _preview_colors(response))
        print(json.dumps({
            "service_uid": os.geteuid(), "private_source_uid": owner.st_uid,
            "private_source_gid": owner.st_gid,
            "service_has_source_group": owner.st_gid in os.getgroups(),
            "normal_embedded_picture_preview": "passed",
            "direct_private_status": direct.status_code,
            "linked_preview_status": response.status_code,
            "private_pixels_in_preview": private_pixels,
            "private_source_unchanged": secret.read_bytes() == png.getvalue(),
        }), flush=True)
        _assert_no_private_picture(response, linked)


def test_real_service_preview_does_not_read_foreign_document_dependencies():
    assert os.geteuid() == 0 and os.environ.get("REQ002_DISPOSABLE_CONTAINER") == "1"
    source_root = str(Path(__file__).resolve().parents[2])
    with tempfile.TemporaryDirectory(prefix="req002-office-service-") as temporary:
        os.chown(temporary, 1000, 1000)
        command = [
            "sudo", "-u", "user", "env", "USER=user", "HOME=/home/user",
            "PYTHONDONTWRITEBYTECODE=1", "PYTHONPATH=" + source_root,
            "TMPDIR=" + temporary,
            "OPEN_TERMINAL_API_KEY=req002-auxiliary-test-not-a-secret",
            "OPEN_TERMINAL_MULTI_USER=true", "OPEN_TERMINAL_USER_PREFIX=rt_",
            "OPEN_TERMINAL_ENABLE_NOTEBOOKS=false", "OPEN_TERMINAL_ENABLE_TERMINAL=true",
            "OPEN_TERMINAL_LOG_DIR=" + str(Path(temporary) / "logs"),
            "OPEN_TERMINAL_PUBLIC_READ_ROOTS=[]",
            "/usr/local/bin/python3-ot", str(Path(__file__).resolve()), "service-worker",
        ]
        result = subprocess.run(command, cwd="/tmp", capture_output=True,
                                text=True, timeout=120)
        assert result.returncode == 0, result.stdout + result.stderr


if __name__ == "__main__":
    assert sys.argv[1:] == ["service-worker"]
    _service_worker()
