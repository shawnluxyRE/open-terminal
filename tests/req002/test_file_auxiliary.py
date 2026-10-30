"""REQ-002 AC-008: helper processes and derived views preserve file ownership.

Comparison and Office extraction use their real workers. Only the optional
LibreOffice-to-PDF converter is isolated, so these tests require no renderer.
"""

import io
import zipfile
from urllib.parse import quote

from docx import Document
import pytest

from test_file_http import denied, put


def test_compare_own_files_runs_real_comparison_worker(api):
    put(api.homes["alice"] / "old.txt", "before\n")
    put(api.homes["alice"] / "new.txt", "after\n")
    response = api.request("POST", "/files/compare", json={"original": "old.txt", "revised": "new.txt"})
    assert response.status_code == 200, response.text
    assert response.json()["additions"] == 1
    assert response.json()["deletions"] == 1
    assert response.json()["original"]["name"] == "old.txt"
    assert response.json()["revised"]["name"] == "new.txt"


@pytest.mark.parametrize("side", ["original", "revised"])
@pytest.mark.parametrize("link", [False, True])
def test_compare_denies_either_private_input_before_worker(api, monkeypatch, side, link):
    from open_terminal.utils import file_compare

    put(api.homes["alice"] / "own.txt", "own\n")
    target = put(api.homes["bob"] / "secret.txt")
    if link:
        path = api.homes["alice"] / "borrowed.txt"
        path.symlink_to(target)
    else:
        path = target
    calls = []

    async def unexpected_worker(*args, **kwargs):
        calls.append((args, kwargs))
        raise AssertionError("A forbidden comparison must not launch a worker")

    monkeypatch.setattr(file_compare, "run_worker", unexpected_worker)
    payload = {"original": "own.txt", "revised": "own.txt", side: str(path)}
    denied(api.request("POST", "/files/compare", json=payload))
    assert calls == []
    assert target.read_text() == "needle private-bob\n"


@pytest.mark.parametrize("forbidden_first", [False, True])
def test_archive_rejects_entire_explicit_list_with_private_input(api, forbidden_first):
    own = put(api.homes["alice"] / "own.txt", "own")
    secret = put(api.homes["bob"] / "secret.txt")
    paths = [str(secret), str(own)] if forbidden_first else [str(own), str(secret)]
    denied(api.request("POST", "/files/archive", json={"paths": paths}))
    assert secret.read_text() == "needle private-bob\n"


def test_public_file_can_be_archived_by_both_users(api):
    path = put(api.public / "manual.txt", "public\n")
    for user in ("alice", "bob"):
        response = api.request("POST", "/files/archive", user=user, json={"paths": [str(path)]})
        assert response.status_code == 200, response.text
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            assert archive.namelist() == ["manual.txt"]
            assert archive.read("manual.txt") == b"public\n"


@pytest.mark.parametrize("entry_type", ["file", "directory"])
def test_archive_supports_owned_unicode_file_and_directory_names(api, entry_type):
    """Normal international filenames remain downloadable through the archive API."""
    home = api.homes["alice"]
    if entry_type == "file":
        source = put(home / "项目说明.txt", "owned unicode filename")
        expected_member = "项目说明.txt"
    else:
        source = home / "项目资料"
        put(source / "说明.txt", "owned unicode filename")
        expected_member = "项目资料/说明.txt"

    response = api.request("POST", "/files/archive", json={"paths": [str(source)]})

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/zip")
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert archive.namelist() == [expected_member]
        assert archive.read(expected_member) == b"owned unicode filename"


def test_office_text_extraction_uses_authorized_own_source(api):
    path = api.homes["alice"] / "notes.docx"
    document = Document()
    document.add_paragraph("Alice document contents")
    document.save(path)
    response = api.request("GET", "/files/read", params={"path": str(path)})
    assert response.status_code == 200, response.text
    assert "Alice document contents" in response.json()["content"]


def test_office_preview_calls_converter_only_for_authorized_source(api, monkeypatch, tmp_path):
    source = api.homes["alice"] / "notes.docx"
    document = Document()
    document.add_paragraph("Own preview")
    document.save(source)
    calls = []
    pdf = b"%PDF-1.4\n% isolated converter output\n%%EOF\n"

    async def converter(path):
        calls.append(path)
        output = tmp_path / "preview.pdf"
        output.write_bytes(pdf)
        return str(output)

    monkeypatch.setattr(api.main, "_preview_pdf", converter)
    response = api.request("GET", "/files/view", params={"path": str(source), "preview": True})
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/pdf")
    assert response.content == pdf
    assert calls == [str(source)]


@pytest.mark.parametrize("endpoint", ["read", "view"])
def test_private_office_link_cannot_reach_extractor_or_preview_fallback(api, monkeypatch, endpoint):
    source = api.homes["bob"] / "secret.docx"
    document = Document()
    document.add_paragraph("Private Bob Office contents")
    document.save(source)
    (api.homes["alice"] / "linked.docx").symlink_to(source)
    calls = []

    async def forbidden_helper(*args, **kwargs):
        calls.append((args, kwargs))
        raise AssertionError("Private Office source reached an external helper")

    monkeypatch.setattr(api.main, "_preview_pdf", forbidden_helper)
    monkeypatch.setattr(api.main, "run_worker", forbidden_helper)
    params = {"path": "linked.docx", "preview": True} if endpoint == "view" else {"path": "linked.docx"}
    denied(api.request("GET", "/files/" + endpoint, params=params))
    assert calls == []


@pytest.mark.parametrize("scope", ["own", "public", "other-user", "linked-other-user"])
def test_static_file_serving_obeys_same_ownership(api, scope):
    root = {"own": api.homes["alice"], "public": api.public,
            "other-user": api.homes["bob"], "linked-other-user": api.homes["bob"]}[scope]
    target = put(root / "page.html", "<p>served fixture</p>")
    if scope == "linked-other-user":
        link = api.homes["alice"] / "linked.html"
        link.symlink_to(target)
        target = link
    response = api.request("GET", "/files/serve/" + quote(str(target).lstrip("/"), safe="/"))
    if scope in ("own", "public"):
        assert response.status_code == 200, response.text
        assert response.content == b"<p>served fixture</p>"
    else:
        denied(response)


def skill(root, name, body):
    folder = root / ".agents/skills" / name
    put(folder / "SKILL.md", f"---\nname: {name}\ndescription: {name} description\n---\n{body}\n")
    return folder


def test_skills_discover_and_read_own_resources_without_private_links(api):
    own = skill(api.homes["alice"], "own", "Alice skill body")
    put(own / "data/own.txt", "own skill resource")
    other = skill(api.homes["bob"], "private", "Bob private skill body")
    secret = put(other / "secret.txt", "Bob private skill resource")
    (own / "forbidden.txt").symlink_to(secret)
    (api.homes["alice"] / ".agents/skills/private-link").symlink_to(other, target_is_directory=True)
    response = api.request("GET", "/skills")
    assert response.status_code == 200, response.text
    assert [entry["name"] for entry in response.json()] == ["own"]
    assert "private" not in response.text
    response = api.request("GET", "/skills/read", params={"name": "own"})
    assert response.status_code == 200, response.text
    assert response.json()["content"] == "Alice skill body"
    assert response.json()["resources"] == ["data/own.txt"]
    response = api.request("GET", "/files/read", params={"path": str(own / "data/own.txt")})
    assert response.status_code == 200, response.text
    assert response.json()["content"] == "own skill resource"
    denied(api.request("GET", "/files/read", params={"path": str(own / "forbidden.txt")}))
    response = api.request("GET", "/skills/read", params={"name": "private"})
    assert response.status_code == 404, response.text
    assert "Bob" not in response.text


def test_whole_skill_root_link_to_other_home_is_excluded(api):
    other = skill(api.homes["bob"], "private", "Bob private skill body")
    own_parent = api.homes["alice"] / ".agents"
    own_parent.mkdir()
    (own_parent / "skills").symlink_to(other.parent, target_is_directory=True)
    response = api.request("GET", "/skills")
    assert response.status_code == 200, response.text
    assert response.json() == []
    response = api.request("GET", "/skills/read", params={"name": "private"})
    assert response.status_code == 404, response.text


def test_skill_with_private_linked_manifest_is_omitted_without_hiding_valid_skills(api):
    skill(api.homes["alice"], "own", "Alice own skill")
    other = skill(api.homes["bob"], "private", "Bob private manifest")
    linked_skill = api.homes["alice"] / ".agents/skills/linked-manifest"
    linked_skill.mkdir()
    (linked_skill / "SKILL.md").symlink_to(other / "SKILL.md")
    response = api.request("GET", "/skills")
    assert response.status_code == 200, response.text
    assert [entry["name"] for entry in response.json()] == ["own"]
    assert "private" not in response.text
    assert "linked-manifest" not in response.text
    response = api.request("GET", "/skills/read", params={"name": "own"})
    assert response.status_code == 200, response.text
    assert response.json()["content"] == "Alice own skill"
    response = api.request("GET", "/skills/read", params={"name": "private"})
    assert response.status_code == 404, response.text
    assert "Bob" not in response.text
